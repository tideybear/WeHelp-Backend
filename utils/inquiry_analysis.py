# -*- coding: utf-8 -*-
# 1. 安装依赖（建议在虚拟环境中进行）
# pip install "paddleocr>=3.0" "paddlepaddle>=3.0"  # CPU版本即可开始
# pip install openai
# pip install pillow pillow-heif   # HEIC(苹果手机照片) 解码需要 pillow-heif

import json
import os
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from paddleocr import PaddleOCR
from PIL import Image, ImageOps

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

try:
    # iPhone 上传的 HEIC 格式需要注册解码器
    import pillow_heif

    pillow_heif.register_heif_opener()
except ImportError:  # pragma: no cover - 没装 pillow-heif 时忽略 HEIC
    pass

from utils.deepseek_analysis import chat  # noqa: E402

# 2. 初始化 OCR 引擎，指定语言为中文+英文
#    注意：PaddleOCR 3.x 中 use_angle_cls 已更名为 use_textline_orientation
ocr = PaddleOCR(use_textline_orientation=True, lang='ch')

# 支持的图片扩展名
IMG_EXTS = {'.jpg', '.jpeg', '.png', '.bmp', '.webp', '.tiff', '.tif', '.gif', '.heic', '.heif'}

# 压缩参数：最长边限制 + JPEG 质量
MAX_SIDE = 1600
JPEG_QUALITY = 85

# 名片字段定义（提示模型按这些字段抽取，可按需增删）
CARD_FIELDS = [
    'name',
    'title',
    'company',
    'country',
    'address',
    'mobile',
    'phone',
    'email',
    'website',
    'core_business',
    'other',
]


# 3. 遍历指定文件夹下的所有图片
def iter_images(folder: str):
    """按文件名排序，遍历 folder 下的所有图片文件（不递归子目录）。"""
    for name in sorted(os.listdir(folder)):
        path = os.path.join(folder, name)
        if os.path.isfile(path) and os.path.splitext(name)[1].lower() in IMG_EXTS:
            yield path


# 3.1 原图压缩：手机上传的原图过大（HEIC 3024x4032 等），统一压缩成合适的 JPG
def compress_to_jpg(src_path: str, out_dir: str, max_side: int = MAX_SIDE,
                    quality: int = JPEG_QUALITY) -> str:
    """把任意格式图片压缩为 JPG，统一输出为 RGB + 最长边不超过 max_side。

    Args:
        src_path: 原图路径
        out_dir: 压缩后图片输出目录
        max_side: 最长边像素上限
        quality: JPEG 质量（1-95，越大越清晰）

    Returns:
        压缩后图片的路径
    """
    os.makedirs(out_dir, exist_ok=True)
    base = os.path.splitext(os.path.basename(src_path))[0]
    out_path = os.path.join(out_dir, f'{base}.jpg')

    with Image.open(src_path) as img:
        # 依据 EXIF 方向旋转，避免手机竖拍图方向错乱
        img = ImageOps.exif_transpose(img)
        # 透明/调色板/灰度图统一转 RGB，否则存 JPG 会失败或偏色
        if img.mode != 'RGB':
            img = img.convert('RGB')

        # 等比缩放到最长边不超过 max_side
        w, h = img.size
        if max(w, h) > max_side:
            scale = max_side / float(max(w, h))
            new_size = (max(1, int(round(w * scale))), max(1, int(round(h * scale))))
            img = img.resize(new_size, Image.LANCZOS)

        img.save(out_path, 'JPEG', quality=quality, optimize=True)

    return out_path


def compress_folder(folder: str, out_dir: str) -> list:
    """把 folder 下所有图片压缩为 JPG 输出到 out_dir，返回 (原图, 压缩图) 列表。

    按原文件名排序，保证后续正反面配对顺序稳定。
    """
    pairs = []
    for src_path in iter_images(folder):
        try:
            out_path = compress_to_jpg(src_path, out_dir)
        except Exception as exc:
            print(f'[压缩失败] {os.path.basename(src_path)}: {exc}')
            continue
        pairs.append((src_path, out_path))
    return pairs


def ocr_image(img_path: str):
    """识别单张图片，返回 [(文本, 置信度), ...]。"""
    items = []
    for res in ocr.predict(img_path):
        texts = res.get('rec_texts') or []
        scores = res.get('rec_scores') or []
        for idx, text in enumerate(texts):
            score = scores[idx] if idx < len(scores) else 0.0
            items.append((text, float(score)))
    return items


def build_prompt(front_texts, back_texts) -> str:
    """把正反面 OCR 文本拼成给模型的抽取提示词。"""
    fields_desc = '\n'.join(f'- {f}' for f in CARD_FIELDS)
    front = '\n'.join(f'- {t}' for t in front_texts) or '(无)'
    back = '\n'.join(f'- {t}' for t in back_texts) or '(无)'
    return f"""你是一名商务名片信息抽取助手。下面是一张名片的正面与背面 OCR 识别文本，
请把它们整理成结构化 JSON。

需要抽取的字段：
{fields_desc}

要求：
1. 只使用 OCR 文本中出现的信息，不要臆测或编造。
2. 某个字段在 OCR 文本中找不到时，值填空字符串 ""。
3. core_business 用简短英文或中文短语概括主营业务（如 electric switch / 金属按钮开关）。
4. website 去掉多余空格与结尾标点。
5. 地址保留原始文字（含邮编）。
6. 只输出 JSON，不要输出任何解释文字或 markdown 代码块。

正面 OCR 文本：
{front}

背面 OCR 文本：
{back}
"""


def parse_json(raw: str) -> dict:
    """从模型回复中解析 JSON，兼容 ```json 代码块与多余文字。"""
    text = raw.strip()
    if '```' in text:
        # 去掉 ```json ... ``` 包裹
        parts = text.split('```')
        for part in parts:
            part = part.strip()
            if part.startswith('json'):
                part = part[4:].strip()
            if part.startswith('{'):
                text = part
                break
    start = text.find('{')
    end = text.rfind('}')
    if start != -1 and end != -1 and end > start:
        text = text[start:end + 1]
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        return {}


def extract_card(front_texts, back_texts) -> dict:
    """调用模型把正反面 OCR 文本整理成结构化名片信息。"""
    raw = chat(build_prompt(front_texts, back_texts), thinking='disabled')
    data = parse_json(raw)
    # 统一字段，缺失补空
    result = {}
    for field in CARD_FIELDS:
        value = data.get(field, '')
        if value is None:
            value = ''
        if not isinstance(value, str):
            value = str(value)
        result[field] = value.strip()
    return result


folder = '/Users/pzj/Downloads/ocrtest'

# Excel 输出文件（与图片同目录，重复运行会追加，不覆盖已有数据）
OUTPUT_EXCEL = os.path.join(folder, 'output.xlsx')

# Excel 列顺序 = CARD_FIELDS + 图片文件名
EXCEL_COLUMNS = list(CARD_FIELDS) + ['front_image', 'back_image']


def append_cards_to_excel(cards, path: str) -> None:
    """把名片数据追加写入 Excel；文件不存在则创建并写表头。"""
    exists = os.path.exists(path)
    if exists:
        wb = load_workbook(path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = 'cards'
        ws.append(EXCEL_COLUMNS)
        # 表头加粗
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center')

    for card in cards:
        ws.append([card.get(col, '') for col in EXCEL_COLUMNS])

    # 列宽自适应
    for col_idx, col_name in enumerate(EXCEL_COLUMNS, start=1):
        width = len(str(col_name).encode('utf-16-le')) // 2 + 4
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx,
                                min_row=1, max_row=ws.max_row):
            value = row[0].value
            if value is not None:
                width = max(width, min(len(str(value).encode('utf-16-le')) // 2 + 2, 60))
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    wb.save(path)


# 4. 先把文件夹下的原图统一压缩成合适的 JPG，再进行识别
COMPRESS_DIR = os.path.join(folder, 'compressed')
pairs = compress_folder(folder, COMPRESS_DIR)
print(f'已压缩 {len(pairs)} 张图片 -> {COMPRESS_DIR}')

# 5. 按名称顺序每两张为一张名片的正反面，整理输出
cards = []
for idx in range(0, len(pairs), 2):
    front_src, front_path = pairs[idx]
    back_src, back_path = pairs[idx + 1] if idx + 1 < len(pairs) else (None, None)

    front_texts = [t for t, _ in ocr_image(front_path)]
    back_texts = [t for t, _ in ocr_image(back_path)] if back_path else []

    card = extract_card(front_texts, back_texts)
    card['front_image'] = os.path.basename(front_src)
    card['back_image'] = os.path.basename(back_src) if back_src else ''

    print(f'===== {card["front_image"]} / {card["back_image"]} =====')
    print(json.dumps(card, ensure_ascii=False, indent=2))
    print()

    cards.append(card)

# 6. 追加写入 Excel
if cards:
    append_cards_to_excel(cards, OUTPUT_EXCEL)
    print(f'已追加 {len(cards)} 条名片记录到: {OUTPUT_EXCEL}')
